from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentos" / "Lista_de_Compras_y_Fallas_Reales_PROJECT_DOMUS.docx"

# compact_reference_guide preset
FONT = "Calibri"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
LIGHT = "F4F6F9"
GRAY = "667085"
GRID = "B8C4D1"
RED = "9B1C1C"
RED_FILL = "FDECEC"
GOLD = "7A5A00"
GOLD_FILL = "FFF4CE"
GREEN = "1F6B4F"
GREEN_FILL = "EAF6F0"
WHITE = "FFFFFF"
BLACK = "222222"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=11, color=BLACK, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    assert sum(widths_dxa) == 9360, widths_dxa
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def paragraph_border_bottom(paragraph, color=BLUE, size="18"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "7")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_node, end])
    set_run_font(run, 8.5, GRAY)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        st = doc.styles[style_name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(size)
        st.font.color.rgb = rgb(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        st = doc.styles[style_name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("PROJECT DOMUS  |  Auditoría de compras")
    set_run_font(r, 8.5, GRAY, bold=True)
    paragraph_border_bottom(p, "D7DBE2", "6")

    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Lista verificada • 1 septiembre 2026  |  Página ")
    set_run_font(r, 8.5, GRAY)
    add_page_field(p)


def add_body(doc, text, bold_prefix=None, italic=False, color=BLACK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, 11, color, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, 11, color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, 11, color, italic=italic)
    return p


def add_callout(doc, label, text, fill=LIGHT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_table_borders(table, accent, "12")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + "\n")
    set_run_font(r, 9, accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, 10.5, BLACK, bold=False)
    set_table_geometry(table, [9360])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths, font_size=9, header_fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, font_size, NAVY, bold=True)
    for row_data in rows:
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if idx in (0,):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, font_size, BLACK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_source(doc, label, url, note):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ": ")
    set_run_font(r, 9.5, NAVY, bold=True)
    add_hyperlink(p, "abrir fuente", url)
    r = p.add_run(" — " + note)
    set_run_font(r, 9.5, BLACK)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    sec = doc.sections[0]

    # First-page pattern: memo_masthead. Preset: compact_reference_guide.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("AUDITORÍA TÉCNICA Y DE COMPRAS")
    set_run_font(r, 10, BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PROJECT DOMUS")
    set_run_font(r, 26, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Lista de compras basada en el inventario real, el firmware actual y precios web en Honduras")
    set_run_font(r, 13, GRAY)
    paragraph_border_bottom(p, BLUE, "18")

    meta = [
        ("Inventario base", "Lista confirmada por Isaac el 1 de septiembre de 2026"),
        ("Tiendas revisadas", "C&D Tecnología y Steren Honduras"),
        ("Destino", "Maqueta escolar de baja tensión; envío a Choluteca por confirmar"),
        ("Criterio", "Comprar solo lo que desbloquea una función demostrable"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + ": ")
        set_run_font(r, 10.5, NAVY, bold=True)
        r = p.add_run(value)
        set_run_font(r, 10.5, BLACK)

    add_callout(
        doc,
        "Veredicto",
        "Sí faltan piezas importantes. No hace falta otra pantalla, pero sí un lector microSD SPI dedicado "
        "si el ESP32 debe abrir archivos, configuraciones, registros o recursos directamente. La ranura del "
        "DFPlayer pertenece al reproductor y no funciona como almacenamiento general del ESP32. "
        "La compra mínima correcta es: fuente estable de 5 V/3 A, módulo de 4 relés, "
        "distribución/fusible/cableado y, si se presentará Jarvis hablando, INMP441 + MAX98357A + altavoz. "
        "El sistema solar debe comprarse después de medir consumo, no antes.",
        GREEN_FILL,
        GREEN,
    )

    doc.add_heading("Presupuesto rápido", level=1)
    add_table(
        doc,
        ["Fase", "Presupuesto", "Qué entrega"],
        [
            ("Casa estable", "hasta L1,020", "5 cargas, alimentación y cableado de demostración"),
            ("Jarvis local + almacenamiento", "L528 + INMP441", "voz local y microSD accesible directamente por el ESP32"),
            ("Respaldo portátil", "L499", "power bank comercial 10,000 mAh; alternativa más segura a celdas sueltas"),
            ("Solar 7,500 mAh", "aprox. L1,355", "fase experimental; sin envío, montaje ni herramientas"),
        ],
        [2500, 1700, 5160],
        9.5,
    )
    add_body(doc, "Los precios son referencias visibles al 1 de septiembre de 2026. En productos con selector, el presupuesto usa la variante necesaria; debe reconfirmarse antes de pagar.", italic=True, color=GRAY, after=0)

    page_break(doc)
    doc.add_heading("1. Fallas reales encontradas", level=1)
    add_callout(doc, "Importante", "No todas las fallas se resuelven comprando. Las filas marcadas como software deben corregirse antes del cableado final.", GOLD_FILL, GOLD)
    blockers = [
        ("CRÍTICA", "GPIO22 usado como SCL", "El ESP32-S3 solo tiene GPIO0–21 y GPIO26–48; el LCD I2C no funcionará así.", "Cambiar pin y validar el pinout físico. No comprar otra pantalla."),
        ("ALTA", "8 relés en código; 1 físico", "Cinco cargas planeadas no pueden conmutarse con un solo canal.", "Comprar 1 módulo de 4 canales y usarlo junto al relé existente."),
        ("ALTA", "Sin barra de 5 V estable", "Bomba, ventilador, relés, audio y LEDs pueden reiniciar el ESP32 por caídas.", "Fuente 5 V/3 A, fusible, cable 22 AWG, conectores y capacitores."),
        ("ALTA", "Jarvis no tiene oído ni voz física", "El firmware principal todavía no puede escuchar ni hablar.", "INMP441, MAX98357A y altavoz; luego integrar y probar modelos."),
        ("ALTA", "Sin almacenamiento microSD del ESP32", "La tarjeta insertada en el DFPlayer solo puede usarse a través de sus órdenes de reproducción; el ESP32 no puede montarla como sistema de archivos.", "Comprar lector SPI separado y una tarjeta FAT32. Reservar la ranura del DFPlayer para audio de respaldo."),
        ("ALTA", "Consumo no medido", "No se puede garantizar batería de 2 h ni tamaño de fusible solo con estimaciones.", "Medir corriente promedio y pico con el multímetro que ya tienes."),
        ("MEDIA", "Sensor de suelo mal descrito", "El hardware confirmado es resistivo, no capacitivo; se corroe y calibra distinto.", "No comprar para la primera demo; corregir texto/código y energizar solo al medir."),
        ("MEDIA", "Sensor de viento inexistente", "El motor con aspa es ventilador, no anemómetro calibrado.", "Eliminar esa afirmación o comprar un sensor real solo si es requisito."),
        ("MEDIA", "PIR y nivel de agua sin lógica", "Los módulos existen, pero hoy no son funciones demostrables.", "No comprar duplicados; asignar pines e integrar software."),
        ("MEDIA", "GPIO19 reservado para USB D−", "El pin del DFPlayer puede interferir si se usa USB nativo.", "Reasignar UART antes de cerrar el arnés."),
        ("MEDIA", "Sin adaptación WS2812", "3.3 V puede funcionar de forma marginal con tira a 5 V.", "74AHCT125/74HCT14 recomendado; el 74HC595 no lo sustituye."),
        ("BAJA", "Capacitor “104 pF” mal rotulado", "104 significa 100 nF = 0.1 µF, no 104 pF.", "No comprar; corregir inventario y usarlo como desacoplo."),
    ]
    add_table(doc, ["Nivel", "Hallazgo", "Consecuencia", "Acción real"], blockers, [850, 2140, 2850, 3520], 8.1)

    page_break(doc)
    doc.add_heading("2. Lo que ya tienes — no duplicar", level=1)
    owned = [
        ("Control", "ESP32-S3 N16R8, Pico, ESP8266", "El ESP32-S3 es el controlador principal; no comprar otra placa."),
        ("Pantalla", "LCD1602 + adaptador I2C", "Suficiente. Confirmar si el adaptador listado dos veces es realmente una o dos unidades."),
        ("Audio legado", "DFPlayer Mini", "Su ranura microSD sirve para las pistas del propio reproductor; no sustituye un lector conectado al ESP32."),
        ("Conmutación", "1 relé de un canal", "Conservarlo; completa cinco cargas junto a un módulo de cuatro canales."),
        ("Carga USB", "2 TP4056 USB-C con protección", "No recomprar. Úsalos solo en una arquitectura USB 1S separada."),
        ("Sensores", "DHT, PIR, HC-SR04, LDR, suelo resistivo, nivel de agua, Reed, RFID", "Hay hardware suficiente; varias funciones faltan por código, no por compras."),
        ("Actuadores", "Bomba, motor/ventilador, servo, LEDs y WS2812", "No comprar duplicados hasta medir consumo y probar cada carga."),
        ("Instrumentación", "Multímetro", "Es la herramienta clave para cerrar batería, fusible y fuente."),
        ("Cableado", "M/M y H/M", "Falta confirmar H/H o decidir soldadura/conectores para el montaje final."),
    ]
    add_table(doc, ["Área", "Confirmado", "Decisión"], owned, [1500, 3100, 4760], 9)

    doc.add_heading("Tres dudas ya resueltas", level=2)
    add_body(doc, "Pantalla: usa el LCD1602 I2C que ya tienes. El soporte OLED del código no crea una necesidad de compra.", bold_prefix="Pantalla:")
    add_body(doc, "microSD: hacen falta un lector SPI separado y una tarjeta para que el ESP32 monte un sistema de archivos. La ranura del DFPlayer solo almacena sus pistas. TinyML puede seguir en la flash del ESP32; la microSD se reserva para datos, configuraciones, registros y recursos que el firmware necesite abrir.", bold_prefix="microSD:")
    add_body(doc, "Relés: compra cuatro canales, no ocho. La maqueta define cinco cargas y ya posees un relé de un canal.", bold_prefix="Relés:")

    page_break(doc)
    doc.add_heading("3. Lista de compras obligatoria — casa estable", level=1)
    add_callout(doc, "Orden recomendado", "Compra esta fase antes de Jarvis y antes del sistema solar. Es la que evita reinicios, cables sobrecargados y funciones que no se pueden demostrar.", GREEN_FILL, GREEN)
    core = [
        ("Fuente 5 V/3 A USB-C con switch", "1", "L350 presup.", "La página muestra desde L248; elegir y confirmar exactamente la variante USB-C 5 V/3 A con interruptor."),
        ("Módulo relé 4 canales 5 V/10 A", "1", "L250 presup.", "Entrada TTL 3.3/5 V, optoacoplado y activo en LOW; completa las cinco cargas."),
        ("Portafusible aéreo 5×20 mm", "1", "L15", "Instalar en la rama principal de 5 V, accesible para reemplazo."),
        ("Fusible 3 A 5×20 mm", "3", "L45", "Uno en uso y dos repuestos. El valor final se confirma al medir picos."),
        ("Capacitores 1000 µF/25 V, paquete de 5", "1", "L35", "Uno en barra principal; otros cerca de audio y cargas con motor."),
        ("Conectores PCT, paquete de 4", "1", "L40", "Derivación desmontable de 5 V y GND; evita empalmes retorcidos."),
        ("Cable cobre flexible 22 AWG", "1 lote", "desde L35", "Alimentación de cargas. No usar jumpers largos para bomba, ventilador o barra principal."),
        ("Jumpers H/H", "1 faja", "hasta L250", "Comprar solo H/H si hay variante; el set C&D trae H/H, H/M y M/M."),
    ]
    add_table(doc, ["Producto", "Cant.", "Precio", "Por qué / condición"], core, [2750, 650, 1150, 4810], 8.7)
    add_body(doc, "Subtotal mínimo conocido sin H/H ni cable: L735. Subtotal conservador con cable y set completo de jumpers: hasta L1,020.", bold_prefix="Subtotal mínimo conocido sin H/H ni cable:")
    add_callout(doc, "Conexión correcta", "La fuente USB-C puede alimentar el ESP32 por su puerto, pero las cargas necesitan una derivación de 5 V propia y GND común. No hagas circular la corriente de bomba, relés y altavoz por el pin de 3.3 V ni por pistas finas del ESP32.", RED_FILL, RED)

    doc.add_heading("Herramientas y acabado eléctrico", level=2)
    add_table(
        doc,
        ["Compra si no existe", "Precio visto", "Uso"],
        [
            ("Termocontráctil, set 127 piezas", "L150", "Aislar soldaduras y cables de la alimentación."),
            ("Estaño 0.8 mm 63/37, 22 ft", "L190", "Montaje final sobre placa o conexiones soldadas."),
            ("Baquelita PCB soldable 7×9 cm", "L80", "Opcional: más robusta que protoboard para exposición."),
        ],
        [3400, 1600, 4360],
        9,
    )

    # Continúa en la misma página si la tabla de herramientas ocupa la página
    # siguiente; evita una página casi vacía antes de la sección 4.
    doc.add_heading("4. Compras para Jarvis local", level=1)
    add_callout(doc, "Límite honesto", "Jarvis será un asistente local de órdenes limitadas: palabra de activación, clasificación de intención y respuestas habladas. No será un LLM conversacional tipo Iron Man.", GOLD_FILL, GOLD)
    voice = [
        ("INMP441 I2S", "1", "Obligatorio", "Sin precio hondureño verificable", "Micrófono digital; alimentar a 3.3 V. No sustituir por KY-037/KY-038."),
        ("MAX98357A", "1", "Obligatorio", "C&D L130", "DAC/amplificador I2S; 2.7–5.5 V y hasta ~3 W en 4 Ω."),
        ("Altavoz 4 Ω/3 W", "1", "Obligatorio", "C&D L90", "Salida hablada; conectar a SPK+ y SPK− del amplificador."),
        ("74AHCT125 o 74HCT14", "1", "Recomendado", "Importación / consultar", "Nivel lógico estable para WS2812 a 5 V."),
        ("Lector/escritor microSD SPI", "1", "Obligatorio", "C&D L149", "Almacenamiento montable por el ESP32. Módulo con SPI y adaptación 3.3/5 V."),
        ("microSD 4 GB FAT32", "1", "Obligatorio", "Usar 32 GB Steren L159 si no hay 4 GB", "Para archivos del ESP32. 2 GB basta, pero 4 GB FAT32 es preferible; 8–32 GB FAT32 también sirve."),
        ("Segunda microSD FAT32", "1", "Opcional", "Steren L159", "Solo si también se usará el DFPlayer como respaldo MP3. Una tarjeta no se comparte entre ambos módulos."),
    ]
    add_table(doc, ["Producto", "Cant.", "Prioridad", "Precio", "Uso real"], voice, [2150, 600, 1350, 1780, 3480], 8.2)
    add_body(doc, "Subtotal de voz y almacenamiento verificable: L528 + precio del INMP441 (MAX98357A, altavoz, lector SPI y una microSD local de 32 GB como sustituta disponible de la tarjeta de 4 GB). Con una segunda tarjeta para el DFPlayer: L687 + INMP441.", bold_prefix="Subtotal de voz y almacenamiento verificable:")
    add_callout(doc, "Dos ranuras, dos funciones", "Lector SPI + tarjeta 1: archivos que abre el ESP32. DFPlayer + tarjeta 2: pistas MP3/WAV que reproduce el DFPlayer. No conectes una misma tarjeta simultáneamente a los dos módulos.", PALE_BLUE, BLUE)
    add_body(doc, "No se encontró una oferta local verificable de un módulo INMP441 en C&D o Steren Honduras. Mouser Honduras muestra una placa oficial de evaluación mucho más cara y sin existencias; para este proyecto conviene buscar el módulo maker INMP441 y confirmar vendedor/envío antes de pagar.")

    doc.add_heading("Lo que falta aunque compres el hardware", level=2)
    for text in (
        "Grabar un dataset español real con ruido de aula y varias voces.",
        "Entrenar y validar el detector de “Jarvis” y el clasificador de intenciones int8.",
        "Integrar PicoTTS con MAX98357A en el firmware principal.",
        "Definir un mapa de pines que no choque con I2C, USB, relés ni sensores.",
        "Probar diez veces seguidas cada orden y mantener desactivado JARVIS_LOCAL_HABILITADO hasta aprobar.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

    # La sección anterior suele terminar al límite de página. Dejar que Word
    # coloque esta sección en la siguiente evita una hoja en blanco intermedia.
    doc.add_heading("5. Energía portátil y solar", level=1)
    add_callout(doc, "Recomendación para la feria", "Para dos horas de demostración, primero mide el consumo. Si quieres movilidad sin fabricar un pack de litio, una power bank comercial de 10,000 mAh y salida 5 V/3 A es la opción más segura y económica encontrada.", GREEN_FILL, GREEN)
    add_table(
        doc,
        ["Opción", "Precio", "Ventaja", "Advertencia"],
        [
            ("Steren MOV-1000, 10,000 mAh", "L499.01", "Protecciones integradas y salida USB hasta 3 A", "Confirmar que no se apague por baja carga y medir autonomía real."),
            ("Pack 1S3P + solar", "aprox. L1,355", "Demuestra carga solar y almacenamiento", "Requiere celdas nuevas, idénticas, igualadas y montaje supervisado."),
        ],
        [2650, 1400, 2500, 2810],
        8.8,
    )

    doc.add_heading("Compra solar condicional — 7,500 mAh nominales", level=2)
    solar = [
        ("Panel 5 V/1,100 mA", "1", "L280", "Debe entregar al menos 4.4 V al CN3065; medir a pleno sol."),
        ("18650 BAK 2,500 mAh", "3", "L540", "Solo nuevas, mismo lote, misma tensión; no mezclar recicladas."),
        ("Portacelda 18650", "3", "L165", "Evita soldar directamente sobre las celdas."),
        ("BMS 1S", "1", "L85", "Una sola protección del pack; no apilar varias placas de protección."),
        ("CN3065 solar", "1", "L135", "Cargador 1S; el IC requiere entrada 4.4–6 V."),
        ("MT3608", "1", "L90", "Eleva 3–4.2 V a 5 V; ajustar con multímetro antes de conectar ESP32."),
        ("Interruptor KCD1", "1", "L25", "Corte general de baja tensión."),
        ("Capacitor 6,800 µF/25 V", "1", "L35", "Ayuda con picos; no reemplaza una fuente o convertidor insuficiente."),
    ]
    add_table(doc, ["Producto", "Cant.", "Subtotal", "Condición técnica"], solar, [2500, 700, 1200, 4960], 8.5)
    add_body(doc, "Subtotal solar estimado: L1,355, sin fusible, cableado, envío ni mano de obra.", bold_prefix="Subtotal solar estimado:")
    add_callout(doc, "No mezclar cargadores", "USB: batería 1S → un TP4056 protegido → elevador. Solar: panel → CN3065 → pack 1S → un BMS → fusible/switch → elevador. No conectes TP4056 + CN3065 + otro BMS en serie “por más protección”.", RED_FILL, RED)
    add_body(doc, "Un panel de 3 V/110 mA no sirve para esta arquitectura: entrega solo 0.33 W y está por debajo de la entrada mínima de 4.4 V del CN3065. El panel de 5 V/1,100 mA también debe medirse, porque la ficha comercial contiene datos de potencia inconsistentes.")

    # La advertencia del panel puede caer al inicio de una página. Mantener la
    # siguiente sección unida al flujo evita dejarla sola en una hoja.
    doc.add_heading("6. Materiales físicos que todavía faltan confirmar", level=1)
    add_body(doc, "Estos artículos no aparecen en el inventario confirmado. Sus precios varían demasiado por tamaño y tienda; se incluyen como lista de compra local sin inventar un total.")
    mechanical = [
        ("Base", "Plywood 6–9 mm o cartón corrugado doble", "Rigidez suficiente para casa, invernadero y panel frontal."),
        ("Muros/techo", "Cartón reciclado, foamboard o plywood delgado", "Elegir un solo sistema estructural y reforzar esquinas."),
        ("Invernadero", "Acrílico/PET transparente reciclado", "Cubierta visible y resistente a salpicaduras."),
        ("Agua", "Depósito cerrado + bandeja de contención", "Evitar fugas sobre electrónica y mesa de exposición."),
        ("Montaje", "Separadores M3, tornillos, tuercas, bridas", "Ninguna placa debe quedar suelta o apoyada sobre metal."),
        ("Adhesivos", "Silicón caliente, pegamento para madera, cinta doble cara", "Fijación mecánica; no usar silicón sobre conectores que deban repararse."),
        ("Acabado", "Pintura negra, azul, tono madera y sellador", "Identidad visual PROJECT DOMUS."),
        ("Etiquetas", "Papel adhesivo o cartulina", "Nombrar sensores, entradas, salidas y riesgos."),
    ]
    add_table(doc, ["Zona", "Comprar / conseguir", "Criterio"], mechanical, [1600, 3500, 4260], 9)

    doc.add_heading("7. No comprar todavía", level=1)
    dont_buy = [
        ("OLED", "El LCD1602 I2C ya cubre la interfaz visual."),
        ("Un segundo lector SPI", "Solo hace falta uno conectado al ESP32; la ranura integrada del DFPlayer ya cubre el audio de respaldo."),
        ("Relé de 8 canales", "Cinco cargas = relé actual de 1 canal + nuevo módulo de 4 canales."),
        ("Otra placa ESP32", "La N16R8 tiene flash/PSRAM suficiente para la arquitectura prevista."),
        ("Otro TP4056", "Ya tienes dos y solo debe usarse uno por arquitectura USB 1S."),
        ("Baterías recicladas mezcladas", "Capacidad, resistencia interna y estado desconocidos elevan el riesgo."),
        ("Panel de 3 V/110 mA", "No llega al mínimo de entrada del CN3065 y su potencia es insuficiente."),
        ("Capacitor enorme como “solución”", "4,700–6,800 µF ayuda a picos, pero no corrige fuente débil ni mal cableado."),
        ("KY-037/KY-038 como sustituto del INMP441", "Son módulos analógicos de sonido y no reemplazan un micrófono I2S para voz."),
    ]
    add_table(doc, ["No comprar", "Razón"], dont_buy, [2800, 6560], 9.2)

    # Si la última fila de “No comprar” pasa de página, la sección 8 debe
    # continuar debajo en vez de crear una hoja huérfana.
    doc.add_heading("8. Orden de compra y pruebas", level=1)
    steps = [
        ("1", "Corregir pinout", "Cambiar GPIO22 y definir pines reales para cinco relés, PIR, agua y audio."),
        ("2", "Medir cargas", "Bomba, ventilador, relés, WS2812 y audio: corriente en reposo, promedio y pico."),
        ("3", "Comprar fase estable", "Fuente, 4 relés, fusible, capacitores, cable y conectores."),
        ("4", "Montar por bloques", "Primero ESP32+LCD; luego relés; después motores; por último audio."),
        ("5", "Comprar Jarvis", "Solo cuando el sistema base pase diez ciclos sin reinicios."),
        ("6", "Decidir autonomía", "Power bank comercial para feria o solar experimental según mediciones."),
        ("7", "Cerrar maqueta", "Después de probar, etiquetar y dejar acceso a fusible, switch y USB."),
    ]
    add_table(doc, ["Paso", "Acción", "Criterio de salida"], steps, [850, 2500, 6010], 9.3)

    doc.add_heading("Checklist antes de pagar", level=2)
    checks = [
        "Confirmar físicamente el modelo exacto y la serigrafía del ESP32-S3.",
        "Enviar a C&D una foto o enlace de la variante: relé 4 canales 5 V/10 A y fuente USB-C 5 V/3 A con switch.",
        "Preguntar existencia, costo y tiempo de envío a Choluteca.",
        "No comprar INMP441 si el vendedor no confirma que es salida digital I2S y alimentación de 3.3 V.",
        "Comprar un lector microSD SPI para el ESP32; confirmar que el módulo acepte lógica de 3.3 V y tarjetas FAT16/FAT32.",
        "Comprar una tarjeta para el lector del ESP32. Si se activará también el DFPlayer, presupuestar una segunda tarjeta independiente.",
        "No comprar celdas para un pack paralelo sin supervisión y sin medir que tengan igual tensión.",
        "Guardar comprobantes y anotar precio real para actualizar el presupuesto final.",
    ]
    for text in checks:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

    doc.add_heading("Aceptación técnica mínima", level=2)
    add_callout(doc, "La maqueta está lista", "Cuando arranca diez veces sin reinicio, activa cada una de las cinco cargas, muestra sensores coherentes en el LCD, corta la bomba por tiempo máximo y mantiene un método manual de respaldo si la voz falla.", GREEN_FILL, GREEN)

    # Mantener el flujo evita que el callout de aceptación quede partido o
    # aislado en una página antes de las fuentes.
    doc.add_heading("9. Fuentes consultadas", level=1)
    add_body(doc, "Enlaces consultados el 1 de septiembre de 2026. Los precios no incluyen envío y pueden cambiar. Las especificaciones técnicas se contrastaron con documentación del fabricante cuando fue posible.", italic=True, color=GRAY)
    sources = [
        ("Espressif — GPIO del ESP32-S3", "https://docs.espressif.com/projects/esp-idf/en/v5.1/esp32s3/api-reference/peripherals/gpio.html", "Confirma GPIO0–21 y GPIO26–48; no existe GPIO22."),
        ("Espressif — DevKitC-1 v1.0", "https://docs.espressif.com/projects/esp-dev-kits/en/latest/esp32s3/esp32-s3-devkitc-1/user_guide_v1.0.html", "Pinout y funciones USB D−/D+ en GPIO19/20."),
        ("C&D — fuente 5 V/3 A", "https://sps.cdtechnologia.net/2985-fuente-para-raspberry-pi3-pi4.html", "Página con variantes micro-USB/USB-C y switch; confirmar precio de variante."),
        ("C&D — relés 1–8 canales", "https://sps.cdtechnologia.net/2719-865-modulo-rele-5-24v-10a-1-8-canales.html", "Incluye variante 4 canales 5 V, TTL 3.3/5 V, optoacoplada y activa en LOW."),
        ("C&D — MAX98357A", "https://sps.cdtechnologia.net/5855-modulo-amplificador-de-audio-i2s-max98357.html", "L130; audio I2S, 2.7–5.5 V, salida para 4 Ω."),
        ("C&D — parlante 4 Ω/3 W", "https://sps.cdtechnologia.net/4094-parlante-4ohm-3w.html", "L90."),
        ("C&D — portafusible", "https://sps.cdtechnologia.net/3685-portafusible-con-cable.html", "L15; 5×20 mm, cable 18 AWG."),
        ("C&D — capacitores", "https://sps.cdtechnologia.net/179-capacitores-electroliticos-1-valor-01-1000uf-5-unidades.html", "Paquete de cinco por L35; elegir 1000 µF."),
        ("C&D — conectores PCT", "https://sps.cdtechnologia.net/2623-conector-de-cable-pct-2-pines-5-pines-4-unidades.html", "Cuatro unidades por L40."),
        ("C&D — jumpers Dupont", "https://sps.cdtechnologia.net/118-1331-jumpers-de-conexion-dunpont-hh-mm-mh-40-hilos-20cm.html", "Selector H/H, H/M y M/M; set mostrado L250."),
        ("C&D — termocontráctil", "https://sps.cdtechnologia.net/productos-rebajados?page=10", "Set de 127 piezas visto a L150."),
        ("Steren — power bank MOV-1000", "https://www.steren.com.hn/power-bank-de-10-000-mah-con-turbo-charge-qc-y-power-delivery-con-2-salidas-usb-y-usb-c.html", "10,000 mAh, salida USB hasta 3 A, L499.01."),
        ("Steren — microSD 32 GB", "https://www.steren.com.hn/memoria-microsd-hc-de-32-gb-clase-uhs-iu1-v10-a1-con-adaptador-sd.html", "L159; alternativa FAT32 disponible cuando no se encuentra una tarjeta de 4 GB."),
        ("C&D — lector/escritor microSD SPI", "https://sps.cdtechnologia.net/509-modulo-lector-y-escritor-de-tarjetas-micro-sd-compatible-con-arduino.html", "L149; SPI, adaptación 3.3/5 V, microSD hasta 2 GB y microSDHC hasta 32 GB."),
        ("DFRobot — DFPlayer Mini", "https://wiki.dfrobot.com/dfr0299", "La tarjeta del DFPlayer almacena las pistas y se controla mediante órdenes UART; no expone un sistema de archivos general al ESP32."),
        ("Espressif — SD por SPI", "https://docs.espressif.com/projects/esp-idf/en/v5.0.3/esp32s3/api-reference/peripherals/sdspi_share.html", "El ESP32-S3 monta una tarjeta conectada directamente por SPI mediante FAT/VFS."),
        ("C&D — panel 3–5.5 V", "https://sps.cdtechnologia.net/4506-panel-solar-3v-55v.html", "Seleccionar 5 V/1,100 mA y verificar datos con multímetro."),
        ("C&D — CN3065", "https://sps.cdtechnologia.net/4804-modulo-de-carga-baterias-de-litio-cn3065.html", "Módulo de carga solar 1S, L135."),
        ("CONSONANCE — CN3065 datasheet", "https://files.seeedstudio.com/wiki/Solar_Charger_Shield_V2.2/res/DSE-CN3065.pdf", "Rango de entrada 4.4–6 V y regulación de batería 4.2 V."),
        ("Adafruit — nivel lógico NeoPixel", "https://learn.adafruit.com/neopixel-levelshifter", "Recomienda 74AHCT125 para traducir 3.3 V a 5 V."),
    ]
    for label, url, note in sources:
        add_source(doc, label, url, note)

    doc.add_heading("10. Conclusión de compra", level=1)
    add_callout(doc, "Compra correcta hoy", "Fase estable (hasta L1,020) + MAX98357A, altavoz, lector microSD SPI y una tarjeta (L528 usando la microSD de 32 GB disponible). El INMP441 se compra solo con vendedor confirmado. OLED y relé de 8 canales no son necesarios. Si se activa también el DFPlayer, hace falta una segunda tarjeta. Para autonomía, prioriza la power bank comercial de L499.01; deja el pack solar para después de medir.", GREEN_FILL, GREEN)
    add_body(doc, "Documento generado a partir del inventario confirmado, la revisión del firmware y fuentes web. No sustituye la supervisión adulta al trabajar con baterías de litio, soldadura o herramientas.", italic=True, color=GRAY, after=0)

    # Core properties
    doc.core_properties.title = "Lista de compras y fallas reales — PROJECT DOMUS"
    doc.core_properties.subject = "Auditoría técnica de componentes faltantes"
    doc.core_properties.author = "Equipo PROJECT DOMUS"
    doc.core_properties.keywords = "ESP32-S3, casa inteligente, compras, Honduras, TinyML, energía solar"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
